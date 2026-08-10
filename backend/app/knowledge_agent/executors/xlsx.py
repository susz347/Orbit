from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.knowledge_agent.executors.base import ChunkDraft, build_chunks
from app.knowledge_agent.models import CorpusProfile, KnowledgeChunk


def _display(value: object) -> str:
    return "" if value is None else str(value).strip()


class SpreadsheetStructuredExecutor:
    strategy_id = "spreadsheet_structured_v1"

    def execute(
        self, source: Path, *, profile: CorpusProfile, run_id: str
    ) -> tuple[KnowledgeChunk, ...]:
        workbook = load_workbook(source, data_only=True, read_only=False)
        drafts: list[ChunkDraft] = []

        try:
            for worksheet in workbook.worksheets:
                rows = list(worksheet.iter_rows(values_only=True))
                nonempty_indexes = [
                    index
                    for index, row in enumerate(rows)
                    if any(_display(value) for value in row)
                ]
                if not nonempty_indexes:
                    continue

                header_index = nonempty_indexes[0]
                header_values = rows[header_index]
                has_stable_header = (
                    sum(bool(_display(value)) for value in header_values) >= 2
                    and len(nonempty_indexes) >= 2
                )
                if has_stable_header:
                    headers = [
                        _display(value) or get_column_letter(index)
                        for index, value in enumerate(header_values, start=1)
                    ]
                    data_start = header_index + 1
                else:
                    headers = [
                        get_column_letter(index)
                        for index in range(1, worksheet.max_column + 1)
                    ]
                    data_start = header_index

                for zero_index, row in enumerate(rows[data_start:], start=data_start):
                    values = [_display(value) for value in row]
                    if not any(values):
                        continue
                    pairs = [
                        f"{headers[index]}: {value}"
                        for index, value in enumerate(values)
                        if value
                    ]
                    row_number = zero_index + 1
                    drafts.append(
                        ChunkDraft(
                            text=" | ".join(pairs),
                            locator=f"sheet:{worksheet.title}:row:{row_number}",
                            sheet=worksheet.title,
                            metadata={
                                "row_number": row_number,
                                "table_quality": profile.table_quality,
                            },
                        )
                    )
        finally:
            workbook.close()

        return build_chunks(
            drafts,
            profile=profile,
            run_id=run_id,
            strategy_id=self.strategy_id,
        )
